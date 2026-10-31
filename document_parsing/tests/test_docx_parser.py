import io

from docx import Document

from document_parsing.docx_parser import parse_docx


def _build_docx(add_headings: bool) -> bytes:
    doc = Document()
    if add_headings:
        doc.add_heading("Title", level=1)
        doc.add_paragraph("Intro text.")
        doc.add_heading("Section One", level=2)
        doc.add_paragraph("Body of section one.")
        doc.add_paragraph("More text.")
        doc.add_heading("Section Two", level=2)
        doc.add_paragraph("Body of section two.")
    else:
        doc.add_paragraph("Just a paragraph of text.")
        doc.add_paragraph("Another paragraph, no headings at all.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


MULTI_HEADING_DOCX = _build_docx(add_headings=True)
NO_HEADINGS_DOCX = _build_docx(add_headings=False)


def test_parse_docx_sections():
    sections = parse_docx(MULTI_HEADING_DOCX)

    assert [s.heading for s in sections] == ["Title", "Section One", "Section Two"]
    assert [s.level for s in sections] == [1, 2, 2]
    assert [s.order for s in sections] == [0, 1, 2]
    assert "Intro text." in sections[0].text
    assert "Body of section one." in sections[1].text
    assert "More text." in sections[1].text
    assert "Body of section two." in sections[2].text


def test_parse_docx_no_headings_fallback():
    sections = parse_docx(NO_HEADINGS_DOCX)

    assert len(sections) == 1
    assert sections[0].heading is None
    assert sections[0].level is None
    assert "Just a paragraph of text." in sections[0].text
    assert "Another paragraph, no headings at all." in sections[0].text


def test_parse_docx_idempotent():
    assert parse_docx(MULTI_HEADING_DOCX) == parse_docx(MULTI_HEADING_DOCX)
