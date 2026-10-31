#!/usr/bin/env python3
"""Generate deterministic, non-sensitive demo/test fixtures under `fixtures/`.

Covers the stacks and source types in docs/HARVESTERS.md: a SQL/dbt repo, a
Java repo, a .NET repo, a Talend-like repo, a sample Postgres database (SQL
fixture), and sample documents including an Excel file. Everything written
here is synthetic; the one "secret" in the Talend fixture is a clearly fake
placeholder used to test future redaction logic.

Usage: python scripts/generate_fixtures.py [output-dir]
Default output dir: <repo-root>/fixtures
"""
from __future__ import annotations

import re
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "fixtures"

FAKE_SECRET = "FAKE_API_KEY=sk-fake-0000000000000000"


def _write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, newline="\n")


def _write_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


def _deterministic_filler(size: int, seed_text: str) -> bytes:
    pattern = seed_text.encode("ascii")
    reps = size // len(pattern) + 1
    return (pattern * reps)[:size]


def _generate_sql_dbt_repo(root: Path) -> None:
    repo = root / "repos" / "sql-dbt-sample"
    _write(
        repo / "dbt_project.yml",
        "name: sample_orders\nversion: '1.0.0'\nprofile: sample_orders\nmodel-paths: ['models']\nseed-paths: ['seeds']\n",
    )
    _write(
        repo / "seeds" / "source_orders.sql",
        "-- Raw source table loaded by the upstream ingestion job (fixture: lineage root)\n"
        "CREATE TABLE source_orders (\n"
        "    order_id BIGINT PRIMARY KEY,\n"
        "    customer_id BIGINT NOT NULL,\n"
        "    order_status TEXT NOT NULL,\n"
        "    order_amount NUMERIC(12, 2) NOT NULL,\n"
        "    created_at TIMESTAMP NOT NULL\n"
        ");\n",
    )
    _write(
        repo / "models" / "staging" / "stg_orders.sql",
        "-- Staging model: cleans source_orders (fixture: lineage middle)\n"
        "SELECT\n"
        "    order_id,\n"
        "    customer_id,\n"
        "    LOWER(order_status) AS order_status,\n"
        "    order_amount,\n"
        "    created_at\n"
        "FROM source_orders\n"
        "WHERE order_amount > 0\n",
    )
    _write(
        repo / "models" / "marts" / "mart_orders.sql",
        "-- Mart model: aggregates stg_orders for the orders dashboard (fixture: lineage leaf)\n"
        "SELECT\n"
        "    customer_id,\n"
        "    order_status,\n"
        "    COUNT(*) AS order_count,\n"
        "    SUM(order_amount) AS total_amount\n"
        "FROM stg_orders\n"
        "GROUP BY customer_id, order_status\n",
    )
    _write(
        repo / "schema.yml",
        "version: 2\n"
        "models:\n"
        "  - name: stg_orders\n"
        "    description: Cleaned orders, one row per order.\n"
        "    columns:\n"
        "      - name: order_id\n"
        "        tests: [unique, not_null]\n"
        "  - name: mart_orders\n"
        "    description: Orders aggregated by customer and status, feeds the orders dashboard.\n",
    )
    _write(
        repo / "README.md",
        "# sample_orders (fixture)\n\nSynthetic dbt-style repo: source_orders -> stg_orders -> mart_orders.\n",
    )


def _generate_java_repo(root: Path) -> None:
    repo = root / "repos" / "java-sample"
    _write(
        repo / "pom.xml",
        "<project>\n"
        "  <modelVersion>4.0.0</modelVersion>\n"
        "  <groupId>com.inspira.sample</groupId>\n"
        "  <artifactId>orders-service</artifactId>\n"
        "  <version>1.0.0</version>\n"
        "</project>\n",
    )
    _write(
        repo / "src" / "main" / "java" / "com" / "inspira" / "sample" / "OrdersRepository.java",
        "package com.inspira.sample;\n\n"
        "// Fixture: reads/writes the sample orders table (lineage: source_orders).\n"
        "public class OrdersRepository {\n"
        "    public void findByCustomer(long customerId) {\n"
        "        // SELECT * FROM source_orders WHERE customer_id = ?\n"
        "    }\n"
        "}\n",
    )
    _write(
        repo / "src" / "main" / "resources" / "orders.sql",
        "-- Fixture: SQL resource bundled with the Java service\n"
        "SELECT order_id, customer_id FROM source_orders WHERE order_status = 'open';\n",
    )
    _write(repo / "README.md", "# orders-service (fixture)\n\nSynthetic Java/Maven service.\n")
    # target/ must be skipped by a harvester per docs/HARVESTERS.md
    _write_bytes(repo / "target" / "classes" / "OrdersRepository.class", b"\xca\xfe\xba\xbe" + b"\x00" * 64)


def _generate_dotnet_repo(root: Path) -> None:
    repo = root / "repos" / "dotnet-sample"
    _write(
        repo / "SampleApp.sln",
        "Microsoft Visual Studio Solution File, Format Version 12.00\n"
        'Project("{FAE04EC0-301F-11D3-BF4B-00C04F79EFBC}") = "SampleApp", "SampleApp.csproj", "{00000000-0000-0000-0000-000000000000}"\n'
        "EndProject\n",
    )
    _write(
        repo / "SampleApp.csproj",
        "<Project Sdk=\"Microsoft.NET.Sdk\">\n"
        "  <PropertyGroup>\n"
        "    <TargetFramework>net8.0</TargetFramework>\n"
        "  </PropertyGroup>\n"
        "</Project>\n",
    )
    _write(
        repo / "OrdersService.cs",
        "namespace SampleApp;\n\n"
        "// Fixture: reads/writes the sample orders table (lineage: source_orders).\n"
        "public class OrdersService\n"
        "{\n"
        "    public void GetOpenOrders()\n"
        "    {\n"
        "        // SELECT * FROM source_orders WHERE order_status = 'open'\n"
        "    }\n"
        "}\n",
    )
    _write(
        repo / "appsettings.json",
        '{\n  "ConnectionStrings": {\n    "Orders": "Host=localhost;Database=sample_orders"\n  }\n}\n',
    )
    # bin/ and obj/ must be skipped by a harvester per docs/HARVESTERS.md
    _write_bytes(repo / "bin" / "Debug" / "SampleApp.dll", b"MZ" + b"\x00" * 64)
    _write(
        repo / "obj" / "SampleApp.csproj.nuget.g.props",
        "<Project ToolsVersion=\"Current\"></Project>\n",
    )


def _generate_talend_repo(root: Path) -> None:
    repo = root / "repos" / "talend-sample"
    _write(
        repo / "talend.project",
        "<project><name>sample_orders_etl</name></project>\n",
    )
    _write(
        repo / "jobs" / "load_orders.item",
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>\n"
        "<!-- Fixture Talend job: loads source_orders into stg_orders -->\n"
        "<job name=\"load_orders\">\n"
        "  <component type=\"tDBInput\" table=\"source_orders\"/>\n"
        "  <component type=\"tDBOutput\" table=\"stg_orders\"/>\n"
        f"  <!-- {FAKE_SECRET} (clearly fake, fixture for redaction tests) -->\n"
        "</job>\n",
    )
    _write(
        repo / "jobs" / "load_orders.properties",
        "job.name=load_orders\nsource.table=source_orders\ntarget.table=stg_orders\n",
    )
    # Oversized dummy files a harvest plan must skip (> 1 MB cap per docs/HARVESTERS.md)
    _write_bytes(
        repo / "screenshots" / "load_orders.png",
        _deterministic_filler(1024 * 1024 + 1024, "FIXTURE-SCREENSHOT-DUMMY-"),
    )
    # Plain-text filler, not a real jar/zip — only its size and extension matter for the
    # size-cap / binary-skip scenario this fixture exercises.
    _write_bytes(
        repo / "lib" / "talend-runtime.jar",
        _deterministic_filler(1024 * 1024 + 2048, "FIXTURE-JAR-DUMMY-"),
    )
    _write(repo / "README.md", "# sample_orders_etl (fixture)\n\nSynthetic Talend-like job repo.\n")


def _generate_shared_scripts(root: Path) -> None:
    base = root / "repos" / "shared" / "scripts"
    _write(
        base / "export_orders.sh",
        "#!/bin/sh\n"
        "# Fixture shell script: reads source_orders, writes a CSV extract.\n"
        "psql -c \"SELECT * FROM source_orders\" > orders_extract.csv\n",
    )
    _write(
        base / "refresh_mart.sh",
        "#!/bin/sh\n"
        "# Fixture shell script: refreshes mart_orders from stg_orders.\n"
        "psql -c \"REFRESH MATERIALIZED VIEW mart_orders\"\n",
    )


def _generate_postgres_fixture(root: Path) -> None:
    sql = """-- Fixture sample Postgres database for harvester tests (KAN-88).
-- Synthetic data only. Mirrors the sql-dbt-sample lineage chain.

CREATE SCHEMA IF NOT EXISTS sample_orders;

CREATE TABLE sample_orders.source_orders (
    order_id BIGINT PRIMARY KEY,
    customer_id BIGINT NOT NULL,
    order_status TEXT NOT NULL,
    order_amount NUMERIC(12, 2) NOT NULL,
    created_at TIMESTAMP NOT NULL
);
COMMENT ON TABLE sample_orders.source_orders IS 'Fixture: raw orders ingested from the upstream order system.';

CREATE TABLE sample_orders.stg_orders (
    order_id BIGINT PRIMARY KEY,
    customer_id BIGINT NOT NULL,
    order_status TEXT NOT NULL,
    order_amount NUMERIC(12, 2) NOT NULL,
    created_at TIMESTAMP NOT NULL
);
COMMENT ON TABLE sample_orders.stg_orders IS 'Fixture: cleaned orders, derived from source_orders.';

CREATE VIEW sample_orders.mart_orders AS
SELECT
    customer_id,
    order_status,
    COUNT(*) AS order_count,
    SUM(order_amount) AS total_amount
FROM sample_orders.stg_orders
GROUP BY customer_id, order_status;
COMMENT ON VIEW sample_orders.mart_orders IS 'Fixture: orders dashboard aggregate, derived from stg_orders.';

CREATE OR REPLACE PROCEDURE sample_orders.refresh_stg_orders()
LANGUAGE SQL
AS $$
    DELETE FROM sample_orders.stg_orders;
    INSERT INTO sample_orders.stg_orders
    SELECT order_id, customer_id, LOWER(order_status), order_amount, created_at
    FROM sample_orders.source_orders
    WHERE order_amount > 0;
$$;
COMMENT ON PROCEDURE sample_orders.refresh_stg_orders() IS 'Fixture: refresh job, mirrors the dbt staging model.';

-- A table deliberately marked stale, for prod-support / broken-job test scenarios.
CREATE TABLE sample_orders.job_runs (
    job_name TEXT PRIMARY KEY,
    last_success_at TIMESTAMP NOT NULL,
    is_stale BOOLEAN NOT NULL DEFAULT FALSE
);
COMMENT ON TABLE sample_orders.job_runs IS 'Fixture: job run log; refresh_stg_orders is marked is_stale=true to simulate a broken job.';

INSERT INTO sample_orders.source_orders (order_id, customer_id, order_status, order_amount, created_at) VALUES
    (1, 100, 'open', 42.50, '2026-01-01 00:00:00'),
    (2, 100, 'closed', 15.00, '2026-01-02 00:00:00'),
    (3, 101, 'open', 99.99, '2026-01-03 00:00:00');

INSERT INTO sample_orders.stg_orders (order_id, customer_id, order_status, order_amount, created_at) VALUES
    (1, 100, 'open', 42.50, '2026-01-01 00:00:00'),
    (2, 100, 'closed', 15.00, '2026-01-02 00:00:00'),
    (3, 101, 'open', 99.99, '2026-01-03 00:00:00');

INSERT INTO sample_orders.job_runs (job_name, last_success_at, is_stale) VALUES
    ('refresh_stg_orders', '2025-01-01 00:00:00', TRUE);
"""
    _write(root / "db" / "init.sql", sql)


def _generate_docs(root: Path) -> None:
    docs = root / "docs"
    _write(
        docs / "data-retention-policy.md",
        "---\n"
        "title: Data Retention Policy\n"
        "sensitivity: restricted\n"
        "---\n\n"
        "# Data Retention Policy (fixture)\n\n"
        "Synthetic policy document. Order records in `source_orders` are retained for 7 years.\n",
    )
    _write(
        docs / "orders-faq.md",
        "---\n"
        "title: Orders FAQ\n"
        "sensitivity: internal\n"
        "---\n\n"
        "# Orders FAQ (fixture)\n\n"
        "**Q: Where does order data come from?**\n"
        "A: The `source_orders` table, refreshed nightly by `refresh_stg_orders`.\n",
    )

    wb = openpyxl.Workbook()
    fixed_time = __import__("datetime").datetime(2026, 1, 1)
    wb.properties.created = fixed_time
    wb.properties.modified = fixed_time
    wb.properties.creator = "context_platform fixtures"
    ws = wb.active
    ws.title = "orders_summary"
    ws.append(["customer_id", "order_status", "order_count", "total_amount"])
    ws.append([100, "open", 1, 42.50])
    ws.append([100, "closed", 1, 15.00])
    ws.append([101, "open", 1, 99.99])
    excel_path = docs / "sample_report.xlsx"
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)
    _normalize_zip_timestamps(excel_path)


def _normalize_zip_timestamps(zip_path: Path) -> None:
    fixed_date_time = (2026, 1, 1, 0, 0, 0)
    tmp_path = zip_path.with_suffix(".tmp")
    with zipfile.ZipFile(zip_path, "r") as src, zipfile.ZipFile(
        tmp_path, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = re.sub(
                    rb"<dcterms:modified[^>]*>[^<]*</dcterms:modified>",
                    b"<dcterms:modified xmlns:dcterms=\"http://purl.org/dc/terms/\" "
                    b'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
                    b'xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:modified>',
                    data,
                )
            item.date_time = fixed_date_time
            dst.writestr(item, data)
    tmp_path.replace(zip_path)


def _generate_expected_md(root: Path) -> None:
    _write(
        root / "EXPECTED.md",
        "# Expected entities and edges per fixture\n\n"
        "Reference for harvester and MCP tests asserting against `fixtures/`. Types per the "
        "Ontology v0 in `CLAUDE.md`.\n\n"
        "## repos/sql-dbt-sample\n"
        "- Entities: Repository(sample_orders), Module(stg_orders), Module(mart_orders), "
        "Table(source_orders)\n"
        "- Edges: DERIVED_FROM(mart_orders -> stg_orders), DERIVED_FROM(stg_orders -> source_orders), "
        "DOCUMENTS(schema.yml -> stg_orders)\n\n"
        "## repos/java-sample\n"
        "- Entities: Repository(orders-service), Module(OrdersRepository), Table(source_orders)\n"
        "- Edges: READS_FROM(OrdersRepository -> source_orders)\n"
        "- `target/` must be excluded from the plan (compiled output).\n\n"
        "## repos/dotnet-sample\n"
        "- Entities: Repository(SampleApp), Module(OrdersService), Table(source_orders)\n"
        "- Edges: READS_FROM(OrdersService -> source_orders)\n"
        "- `bin/`, `obj/` must be excluded from the plan (build output).\n\n"
        "## repos/talend-sample\n"
        "- Entities: Repository(sample_orders_etl), Job(load_orders), Table(source_orders), "
        "Table(stg_orders)\n"
        "- Edges: READS_FROM(load_orders -> source_orders), WRITES_TO(load_orders -> stg_orders)\n"
        "- `screenshots/`, `lib/` must be excluded from the plan (over the 1 MB size cap / binaries).\n"
        "- `load_orders.item` contains one deliberately fake secret for redaction tests; it must be "
        "redacted before storage and flagged with a non-default `sensitivity`.\n\n"
        "## db/init.sql\n"
        "- Entities: Table(source_orders), Table(stg_orders), Dashboard(mart_orders), "
        "Job(refresh_stg_orders)\n"
        "- Edges: DERIVED_FROM(stg_orders -> source_orders), READS_FROM(mart_orders -> stg_orders), "
        "WRITES_TO(refresh_stg_orders -> stg_orders)\n"
        "- `job_runs.refresh_stg_orders.is_stale = true` simulates a broken upstream job.\n\n"
        "## docs\n"
        "- Entities: Document(data-retention-policy.md, sensitivity=restricted), "
        "Document(orders-faq.md, sensitivity=internal), Document(sample_report.xlsx)\n"
        "- Edges: DOCUMENTS(orders-faq.md -> source_orders), "
        "DOCUMENTS(data-retention-policy.md -> source_orders)\n",
    )


def generate(output_dir: Path) -> None:
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True)
    _generate_sql_dbt_repo(output_dir)
    _generate_java_repo(output_dir)
    _generate_dotnet_repo(output_dir)
    _generate_talend_repo(output_dir)
    _generate_shared_scripts(output_dir)
    _generate_postgres_fixture(output_dir)
    _generate_docs(output_dir)
    _generate_expected_md(output_dir)


def main() -> int:
    output_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUTPUT
    generate(output_dir)
    print(f"Fixtures written to {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
